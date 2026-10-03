import fitz  # PyMuPDF
import logging
import io
import csv
from docx import Document

logger = logging.getLogger(__name__)

MAX_CHARS = 30000


def _truncate(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + "\n\n[Text truncated because document is too large]"
    return text


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts = []

        for page_num, page in enumerate(doc, start=1):
            page_text = page.get_text("text")
            if page_text and page_text.strip():
                text_parts.append(f"--- Page {page_num} ---\n{page_text.strip()}")

        doc.close()
        return _truncate("\n\n".join(text_parts))
    except Exception as e:
        logger.exception("Failed to extract text from PDF.")
        raise RuntimeError(f"PDF_EXTRACTION_ERROR: {str(e)}") from e


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        text_parts = []

        for para in doc.paragraphs:
            if para.text and para.text.strip():
                text_parts.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))

        return _truncate("\n\n".join(text_parts))
    except Exception as e:
        logger.exception("Failed to extract text from DOCX.")
        raise RuntimeError(f"DOCX_EXTRACTION_ERROR: {str(e)}") from e


def extract_text_from_plain(file_bytes: bytes) -> str:
    try:
        # Try common encodings
        for encoding in ("utf-8", "utf-16", "latin-1", "cp1252"):
            try:
                text = file_bytes.decode(encoding)
                return _truncate(text)
            except UnicodeDecodeError:
                continue
        raise RuntimeError("Could not decode text file")
    except Exception as e:
        logger.exception("Failed to extract text from plain file.")
        raise RuntimeError(f"TEXT_EXTRACTION_ERROR: {str(e)}") from e


def extract_text_from_csv(file_bytes: bytes) -> str:
    try:
        text = None
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                text = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            raise RuntimeError("Could not decode CSV file")

        reader = csv.reader(io.StringIO(text))
        rows = []
        for i, row in enumerate(reader):
            if i >= 500:  # limit rows
                rows.append("... [more rows truncated]")
                break
            rows.append(" | ".join(cell.strip() for cell in row))

        return _truncate("\n".join(rows))
    except Exception as e:
        logger.exception("Failed to extract text from CSV.")
        raise RuntimeError(f"CSV_EXTRACTION_ERROR: {str(e)}") from e


def extract_text_from_document(file_bytes: bytes, filename: str) -> str:
    filename = (filename or "").lower()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)

    if filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)

    if filename.endswith((".txt", ".md", ".markdown")):
        return extract_text_from_plain(file_bytes)

    if filename.endswith(".csv"):
        return extract_text_from_csv(file_bytes)

    raise RuntimeError("UNSUPPORTED_FILE_TYPE")